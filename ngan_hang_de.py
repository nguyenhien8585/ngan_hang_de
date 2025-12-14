"""
NGÂN HÀNG CÂU HỎI & TẠO ĐỀ THI v3.5 ULTIMATE - FIX HOÀN TẤT
Tự động nhận diện 4 loại câu hỏi + Import/Export Database + Excel tổng hợp
FIX: Tương thích với database cũ có RGBColor lỗi
"""

import customtkinter as ctk
from tkinter import filedialog, messagebox, BooleanVar
import docx
from docx.shared import Pt, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import random
import os
import io
import copy
import threading
import re
from lxml import etree
import traceback
import sqlite3
import pickle
import shutil


# ===================== HELPER FUNCTIONS =====================

def safe_normalize_rgb(color_value):
    """Chuyển đổi bất kỳ giá trị color nào thành tuple (r,g,b) an toàn"""
    if color_value is None:
        return None

    try:
        # Nếu đã là tuple/list hợp lệ
        if isinstance(color_value, (tuple, list)) and len(color_value) >= 3:
            r = max(0, min(255, int(color_value[0]) if color_value[0] is not None else 0))
            g = max(0, min(255, int(color_value[1]) if color_value[1] is not None else 0))
            b = max(0, min(255, int(color_value[2]) if color_value[2] is not None else 0))
            return (r, g, b)

        # Nếu có __iter__ (RGBColor hoặc object tương tự)
        if hasattr(color_value, '__iter__'):
            try:
                rgb_list = list(color_value)[:3]
                r = max(0, min(255, int(rgb_list[0]) if rgb_list[0] is not None else 0))
                g = max(0, min(255, int(rgb_list[1]) if rgb_list[1] is not None else 0))
                b = max(0, min(255, int(rgb_list[2]) if rgb_list[2] is not None else 0))
                return (r, g, b)
            except:
                pass

        # Nếu là int (giả sử là giá trị hex hoặc single value)
        if isinstance(color_value, int):
            return (0, 0, 0)  # Default to black

        return None
    except:
        return None


def repair_content_element(element):
    """Sửa chữa ContentElement từ database cũ"""
    if not hasattr(element, 'formatting') or not element.formatting:
        element.formatting = {}
        return element

    # Fix font_color nếu có
    if 'font_color' in element.formatting:
        old_color = element.formatting['font_color']
        element.formatting['font_color'] = safe_normalize_rgb(old_color)

    return element


def repair_question(question):
    """Sửa chữa Question object từ database cũ"""
    try:
        # Repair question_body
        if hasattr(question, 'question_body'):
            question.question_body = [repair_content_element(el) for el in question.question_body]

        # Repair options
        if hasattr(question, 'options'):
            for option in question.options:
                if hasattr(option, 'content'):
                    option.content = [repair_content_element(el) for el in option.content]

        # Repair solution_body
        if hasattr(question, 'solution_body'):
            question.solution_body = [repair_content_element(el) for el in question.solution_body]

        return question
    except Exception as e:
        print(f'Lỗi repair question: {e}')
        return question


# ===================== DATA MODELS =====================

class ContentElement:
    """Lưu trữ một phần nội dung (văn bản, hình ảnh, hoặc công thức)"""

    def __init__(self, type, data, formatting=None):
        self.type = type  # 'text', 'image', 'equation'
        self.data = data

        # FIX: Chuẩn hóa formatting ngay khi khởi tạo
        if formatting and isinstance(formatting, dict):
            self.formatting = self._normalize_formatting(formatting)
        else:
            self.formatting = formatting or {}

        self.width = None
        self.height = None
        if type == 'image' and self.formatting:
            self.width = self.formatting.get('width')
            self.height = self.formatting.get('height')

    def _normalize_formatting(self, fmt):
        """Chuẩn hóa formatting để tránh lỗi RGBColor"""
        normalized = {}

        for key, value in fmt.items():
            if key == 'font_color':
                normalized[key] = safe_normalize_rgb(value)
            else:
                normalized[key] = value

        return normalized


class Option:
    """Lưu trữ một phương án lựa chọn"""

    def __init__(self, original_label):
        self.original_label = original_label
        self.content = []
        self.is_fixed = False


class Question:
    """Lưu trữ một câu hỏi hoàn chỉnh"""

    def __init__(self, type, original_number_text):
        self.type = type
        self.original_number_text = original_number_text
        self.question_body = []
        self.options = []
        self.solution_body = []
        self.correct_answer = ''
        self.new_number_text = ''
        self.new_correct_answer = ''


class Part:
    """Lưu trữ một phần của đề thi"""

    def __init__(self, type, title_elements):
        self.type = type
        self.title_elements = title_elements
        self.intro_elements = []
        self.questions = []


# ===================== DATABASE MANAGER =====================

class DatabaseManager:
    """Quản lý cơ sở dữ liệu câu hỏi"""

    DB_NAME = 'question_bank.db'

    @staticmethod
    def init_db():
        """Khởi tạo database"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        cursor.execute('''
        CREATE TABLE IF NOT EXISTS Topics (
            topic_id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL
        )
        ''')

        cursor.execute('''
        CREATE TABLE IF NOT EXISTS Questions (
            question_id INTEGER PRIMARY KEY AUTOINCREMENT,
            topic_id INTEGER NOT NULL,
            question_object BLOB NOT NULL,
            question_type INTEGER NOT NULL,
            FOREIGN KEY (topic_id) REFERENCES Topics (topic_id)
        )
        ''')

        conn.commit()
        conn.close()

    @staticmethod
    def export_database(export_path):
        """Export database ra file"""
        try:
            shutil.copy2(DatabaseManager.DB_NAME, export_path)
            return True
        except Exception as e:
            print(f'Lỗi export: {e}')
            return False

    @staticmethod
    def import_database(import_path):
        """Import database từ file"""
        try:
            if os.path.exists(DatabaseManager.DB_NAME):
                backup_path = DatabaseManager.DB_NAME + '.backup'
                shutil.copy2(DatabaseManager.DB_NAME, backup_path)

            shutil.copy2(import_path, DatabaseManager.DB_NAME)
            return True
        except Exception as e:
            print(f'Lỗi import: {e}')
            return False

    @staticmethod
    def merge_database(import_path):
        """Merge database từ file khác vào database hiện tại"""
        try:
            conn_current = sqlite3.connect(DatabaseManager.DB_NAME)
            cursor_current = conn_current.cursor()

            cursor_current.execute(f"ATTACH DATABASE '{import_path}' AS import_db")

            cursor_current.execute("SELECT topic_id, name FROM import_db.Topics")
            import_topics = cursor_current.fetchall()

            topic_mapping = {}

            for old_topic_id, topic_name in import_topics:
                cursor_current.execute("SELECT topic_id FROM Topics WHERE name = ?", (topic_name,))
                result = cursor_current.fetchone()

                if result:
                    new_topic_id = result[0]
                else:
                    cursor_current.execute("INSERT INTO Topics (name) VALUES (?)", (topic_name,))
                    new_topic_id = cursor_current.lastrowid

                topic_mapping[old_topic_id] = new_topic_id

            cursor_current.execute("""
                SELECT question_id, topic_id, question_object, question_type 
                FROM import_db.Questions
            """)
            import_questions = cursor_current.fetchall()

            for _, old_topic_id, question_object, question_type in import_questions:
                new_topic_id = topic_mapping[old_topic_id]
                cursor_current.execute("""
                    INSERT INTO Questions (topic_id, question_object, question_type)
                    VALUES (?, ?, ?)
                """, (new_topic_id, question_object, question_type))

            cursor_current.execute("DETACH DATABASE import_db")
            conn_current.commit()
            conn_current.close()

            return True
        except Exception as e:
            print(f'Lỗi merge: {e}')
            traceback.print_exc()
            return False

    @staticmethod
    def get_or_create_topic(topic_name):
        """Lấy hoặc tạo mới topic"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        cursor.execute('SELECT topic_id FROM Topics WHERE name = ?', (topic_name,))
        result = cursor.fetchone()

        if result:
            topic_id = result[0]
        else:
            cursor.execute('INSERT INTO Topics (name) VALUES (?)', (topic_name,))
            topic_id = cursor.lastrowid
            conn.commit()

        conn.close()
        return topic_id

    @staticmethod
    def insert_question(topic_id, question_type, question_blob):
        """Thêm câu hỏi vào database"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        cursor.execute(
            'INSERT INTO Questions (topic_id, question_object, question_type) VALUES (?, ?, ?)',
            (topic_id, question_blob, question_type)
        )

        conn.commit()
        conn.close()

    @staticmethod
    def get_all_topics():
        """Lấy danh sách tất cả topics"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        cursor.execute('SELECT topic_id, name FROM Topics')
        topics = cursor.fetchall()

        conn.close()
        return topics

    @staticmethod
    def get_topic_name(topic_id):
        """Lấy tên topic"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        cursor.execute('SELECT name FROM Topics WHERE topic_id = ?', (topic_id,))
        result = cursor.fetchone()

        conn.close()
        return result[0] if result else f"Topic {topic_id}"

    @staticmethod
    def count_questions_by_topic_and_type(topic_id, q_type):
        """Đếm số câu hỏi theo topic và loại"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        cursor.execute('''
            SELECT COUNT(question_id) FROM Questions 
            WHERE topic_id = ? AND question_type = ?
        ''', (topic_id, q_type))

        count = cursor.fetchone()[0]
        conn.close()
        return count

    @staticmethod
    def get_random_questions(topic_id, q_type, limit):
        """Lấy ngẫu nhiên câu hỏi - VỚI REPAIR CHO DATABASE CŨ"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        query = '''
            SELECT question_object FROM Questions 
            WHERE topic_id = ? AND question_type = ?
            ORDER BY RANDOM()
            LIMIT ?
        '''

        cursor.execute(query, (topic_id, q_type, limit))
        rows = cursor.fetchall()

        conn.close()

        # CRITICAL FIX: Repair questions khi load từ database
        repaired_blobs = []
        for row in rows:
            try:
                q = pickle.loads(row[0])
                q = repair_question(q)  # Sửa chữa question
                repaired_blobs.append(pickle.dumps(q))  # Re-serialize
            except Exception as e:
                print(f'Lỗi repair question: {e}')
                repaired_blobs.append(row[0])  # Giữ nguyên nếu lỗi

        return repaired_blobs

    @staticmethod
    def delete_topic_and_questions(topic_id):
        """Xóa topic và tất cả câu hỏi liên quan"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        try:
            cursor.execute('DELETE FROM Questions WHERE topic_id = ?', (topic_id,))
            cursor.execute('DELETE FROM Topics WHERE topic_id = ?', (topic_id,))
            conn.commit()
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()

    @staticmethod
    def clear_all_questions():
        """Xóa toàn bộ database"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        try:
            cursor.execute('DELETE FROM Questions')
            cursor.execute('DELETE FROM Topics')
            conn.commit()
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()

    @staticmethod
    def rename_topic(topic_id, new_name):
        """Đổi tên topic"""
        conn = sqlite3.connect(DatabaseManager.DB_NAME)
        cursor = conn.cursor()

        try:
            cursor.execute('UPDATE Topics SET name = ? WHERE topic_id = ?', (new_name, topic_id))
            conn.commit()
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()


# ===================== WORD PARSER =====================

def natural_sort_key(s):
    """Sắp xếp tự nhiên"""
    if not isinstance(s, str):
        return s
    return [int(text) if text.isdigit() else text.lower() for text in re.split('([0-9]+)', s)]


def get_equation_xml(run):
    """Trích xuất XML của công thức toán học (Equation/MathType)"""
    try:
        xml_str = run.element.xml
        root = etree.fromstring(xml_str)

        MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

        omath = root.find(f'.//{{{MATH_NS}}}oMath')
        if omath is not None:
            return etree.tostring(omath, encoding='unicode')

        for obj_elem in root.findall('.//w:object',
                                     namespaces={'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}):
            return etree.tostring(obj_elem, encoding='unicode')

        return None
    except Exception as e:
        print(f'Lỗi trích xuất công thức: {e}')
        return None


def get_image_blob(p, run):
    """Trích xuất blob hình ảnh từ run"""
    try:
        xml_str = run.element.xml
        root = etree.fromstring(xml_str)
        PIC_NAMESPACE = 'http://schemas.openxmlformats.org/drawingml/2006/picture'
        A_NAMESPACE = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        R_NAMESPACE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        WP_NAMESPACE = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'

        pic = root.find(f'.//{{{PIC_NAMESPACE}}}pic')
        if pic is None:
            return (None, None)

        blip = pic.find(f'.//{{{A_NAMESPACE}}}blip')
        if blip is None:
            return (None, None)

        embed = blip.get(f'{{{R_NAMESPACE}}}embed')
        if not embed:
            return (None, None)

        image_part = p.part.related_parts.get(embed)
        if not image_part:
            return (None, None)

        image_blob = image_part.blob

        width = None
        height = None

        extent = root.find(f'.//{{{WP_NAMESPACE}}}extent')
        if extent is not None:
            cx = extent.get('cx')
            cy = extent.get('cy')
            if cx and cy:
                width = int(cx)
                height = int(cy)

        if width is None or height is None:
            xfrm = root.find(f'.//{{{A_NAMESPACE}}}xfrm')
            if xfrm is not None:
                ext = xfrm.find(f'{{{A_NAMESPACE}}}ext')
                if ext is not None:
                    cx = ext.get('cx')
                    cy = ext.get('cy')
                    if cx and cy:
                        width = int(cx)
                        height = int(cy)

        return (image_blob, {'width': width, 'height': height})
    except Exception as e:
        print(f'Lỗi trích xuất ảnh: {e}')
        return (None, None)


def parse_paragraph_elements(p):
    """Parse các element trong paragraph (text, image, equation)"""
    elements = []

    for run in p.runs:
        equation_xml = get_equation_xml(run)
        if equation_xml:
            elements.append(ContentElement('equation', equation_xml))
            continue

        if '<w:drawing>' in run.element.xml or '<w:pict>' in run.element.xml:
            blob, dims = get_image_blob(p, run)
            if blob:
                elements.append(ContentElement('image', blob, formatting=dims))
                continue

        if not run.text:
            continue

        font_size_pt = None
        if run.font.size is not None:
            font_size_pt = run.font.size.pt

        # Chuyển font_color thành tuple ngay lập tức
        font_color_tuple = None
        try:
            if run.font.color and run.font.color.rgb:
                font_color_tuple = safe_normalize_rgb(run.font.color.rgb)
        except:
            font_color_tuple = None

        formatting = {
            'bold': run.bold,
            'italic': run.italic,
            'underline': run.underline,
            'font_name': run.font.name,
            'font_size': font_size_pt,
            'font_color': font_color_tuple
        }
        elements.append(ContentElement('text', run.text, formatting))

    if not p.runs:
        elements.append(ContentElement('text', '\n', {}))
    else:
        elements.append(ContentElement('text', '\n', {}))

    return elements


def detect_question_type(p_text, elements, next_paragraphs_text):
    """TỰ ĐỘNG NHẬN DIỆN LOẠI CÂU HỎI"""
    if re.match(r'^Bài\s+\d+', p_text, re.IGNORECASE):
        return 4

    combined_text = p_text + '\n' + '\n'.join(next_paragraphs_text[:10])

    has_abcd_options = bool(re.search(r'\b[A-D]\.\s+', combined_text))
    has_abcd_parenthesis = bool(re.search(r'\b[a-d]\)\s+', combined_text))

    if has_abcd_options:
        return 1
    elif has_abcd_parenthesis:
        return 2
    elif re.match(r'^Câu\s+\d+', p_text, re.IGNORECASE):
        return 3
    else:
        return 3


def parse_word_doc(filepath):
    """Parse file Word thành các Part và Question - TỰ ĐỘNG NHẬN DIỆN"""
    try:
        doc = docx.Document(filepath)

        all_paragraphs_text = [p.text.strip() for p in doc.paragraphs]

        questions = []
        current_question = None
        current_option = None
        state = 'START'

        for idx, p in enumerate(doc.paragraphs):
            p_text = p.text.strip()
            elements = parse_paragraph_elements(p)

            next_paragraphs_text = all_paragraphs_text[idx + 1:idx + 15]

            if re.match(r'^(Câu|Bài)\s*\d+', p_text, re.IGNORECASE):
                if current_question:
                    questions.append(current_question)

                q_type = detect_question_type(p_text, elements, next_paragraphs_text)

                current_question = Question(
                    type=q_type,
                    original_number_text=p_text.split('.')[0]
                )
                current_question.question_body.extend(elements)
                state = 'QUESTION'
                current_option = None

            elif state in ['QUESTION', 'OPTION'] and current_question and current_question.type == 1:
                if re.match(r'^[A-D]\.', p_text):
                    label = p_text[0]
                    current_option = Option(original_label=label)
                    if re.match(r'^[A-D]\.\s*\*', p_text):
                        current_option.is_fixed = True
                    current_question.options.append(current_option)
                    current_option.content.extend(elements)
                    state = 'OPTION'
                elif current_option:
                    if current_option.content:
                        current_option.content.append(ContentElement('text', '\n', {}))
                    current_option.content.extend(elements)
                elif current_question:
                    if current_question.question_body:
                        current_question.question_body.append(ContentElement('text', '\n', {}))
                    current_question.question_body.extend(elements)

            elif state in ['QUESTION', 'OPTION'] and current_question and current_question.type == 2:
                if re.match(r'^[a-d]\)', p_text):
                    label = p_text[0]
                    current_option = Option(original_label=label)
                    current_question.options.append(current_option)
                    current_option.content.extend(elements)
                    state = 'OPTION'
                elif current_option:
                    if current_option.content:
                        current_option.content.append(ContentElement('text', '\n', {}))
                    current_option.content.extend(elements)
                elif current_question:
                    if current_question.question_body:
                        current_question.question_body.append(ContentElement('text', '\n', {}))
                    current_question.question_body.extend(elements)

            elif p_text.startswith('Lời giải:'):
                if current_question:
                    current_question.solution_body.extend(elements)
                    state = 'SOLUTION'

            elif p_text.startswith('Đáp án đúng:'):
                if current_question and current_question.type != 4:
                    current_question.correct_answer = p_text.split(':', 1)[1].strip()
                state = 'ANSWER'

            else:
                if not elements:
                    continue

                if elements[0].type == 'text' and elements[0].data == '\n':
                    if state == 'QUESTION' and current_question:
                        current_question.question_body.extend(elements)
                    continue

                if state == 'QUESTION' and current_question:
                    if current_question.question_body:
                        current_question.question_body.append(ContentElement('text', '\n', {}))
                    current_question.question_body.extend(elements)
                elif state == 'OPTION' and current_option:
                    if current_option.content:
                        current_option.content.append(ContentElement('text', '\n', {}))
                    current_option.content.extend(elements)
                elif state == 'SOLUTION' and current_question:
                    if current_question.solution_body:
                        current_question.solution_body.append(ContentElement('text', '\n', {}))
                    current_question.solution_body.extend(elements)

        if current_question:
            questions.append(current_question)

        parts = []
        parts_dict = {1: [], 2: [], 3: [], 4: []}

        for q in questions:
            parts_dict[q.type].append(q)

        type_names = {
            1: 'Trắc nghiệm (ABCD)',
            2: 'Đúng/Sai (abcd)',
            3: 'Trả lời ngắn',
            4: 'Tự luận'
        }

        for q_type in [1, 2, 3, 4]:
            if parts_dict[q_type]:
                part = Part(
                    type=q_type,
                    title_elements=[
                        ContentElement('text', f'Phần {q_type}: {type_names[q_type]}', {'bold': True, 'font_size': 14})]
                )
                part.questions = parts_dict[q_type]
                parts.append(part)

        return parts

    except Exception as e:
        raise Exception(f'Không thể parse file Word.\nLỗi: {e}')


# ===================== SCRAMBLE LOGIC =====================

def scramble_data(original_parts, scramble_parts_flag, scramble_questions_flag,
                  scramble_mcq_options_flag, scramble_tf_options_flag):
    """Xáo trộn đề thi"""
    scrambled_parts = copy.deepcopy(original_parts)
    answer_key = {}

    if scramble_parts_flag:
        random.shuffle(scrambled_parts)

    for part in scrambled_parts:
        if scramble_questions_flag:
            if part.type in [1, 2, 3]:
                random.shuffle(part.questions)

        part_key = []
        part_title_texts = []

        if part.title_elements:
            for el in part.title_elements:
                if el.type == 'text' and el.data != '\n':
                    part_title_texts.append(el.data)

        part_title = ''.join(part_title_texts).strip()
        if not part_title:
            part_title = f'PHẦN {part.type}'

        for new_idx, q in enumerate(part.questions):
            if part.type == 4:
                q.new_number_text = f'Bài {new_idx + 1}'

            try:
                if part.type == 1:
                    correct_opt = next(opt for opt in q.options if opt.original_label == q.correct_answer)

                    if scramble_mcq_options_flag:
                        scramblable_options = [opt for opt in q.options if not opt.is_fixed]
                        fixed_options = [opt for opt in q.options if opt.is_fixed]
                        random.shuffle(scramblable_options)
                        q.options = scramblable_options + fixed_options

                    new_labels = ['A', 'B', 'C', 'D']
                    # Ensure we don't exceed available labels
                    if len(q.options) > len(new_labels):
                        q.options = q.options[:len(new_labels)]
                    
                    new_correct_idx = q.options.index(correct_opt)
                    if new_correct_idx < len(new_labels):
                        q.new_correct_answer = new_labels[new_correct_idx]
                    else:
                        q.new_correct_answer = q.correct_answer
                    part_key.append(f'{new_idx + 1}. {q.new_correct_answer}')

                elif part.type == 2:
                    if scramble_tf_options_flag and q.options:
                        try:
                            cleaned_answer = q.correct_answer.strip()
                            original_answers = re.split(r'[;,\s]+', cleaned_answer)
                            original_answers = [ans for ans in original_answers if ans]

                            if len(original_answers) != len(q.options):
                                raise Exception(
                                    f'Lỗi dữ liệu: {q.original_number_text} có {len(q.options)} phương án '
                                    f'nhưng có {len(original_answers)} đáp án.'
                                )

                            new_labels = ['a', 'b', 'c', 'd']
                            # Limit to available labels
                            num_options = min(len(q.options), len(new_labels))
                            old_to_new = {}

                            temp_options = q.options[:num_options]
                            random.shuffle(temp_options)
                            q.options = temp_options

                            for i in range(len(temp_options)):
                                if i < len(new_labels):
                                    opt = temp_options[i]
                                    old_to_new[opt.original_label] = new_labels[i]

                            new_answers = [old_to_new.get(ans, ans) for ans in original_answers]
                            q.new_correct_answer = '; '.join(new_answers)
                        except Exception as e:
                            print(f'Lỗi xáo trộn: {e}')
                            q.new_correct_answer = q.correct_answer
                    else:
                        q.new_correct_answer = q.correct_answer

                    part_key.append(f'{new_idx + 1}. {q.new_correct_answer}')

                elif part.type == 3:
                    q.new_correct_answer = q.correct_answer
                    part_key.append(f'{new_idx + 1}. {q.new_correct_answer}')

                elif part.type == 4:
                    q.new_correct_answer = '(Tự luận)'
                    part_key.append(f'{new_idx + 1}. {q.new_correct_answer}')

            except StopIteration:
                print(f'Không tìm thấy đáp án: {q.original_number_text}')
                q.new_correct_answer = q.correct_answer
            except Exception as e:
                print(f'Lỗi xử lý câu {q.original_number_text}: {e}')
                q.new_correct_answer = q.correct_answer

        answer_key[part_title] = part_key

    return (scrambled_parts, answer_key)


# ===================== WORD WRITER =====================

def insert_equation_xml(paragraph, equation_xml):
    """Chèn công thức toán học vào paragraph"""
    try:
        equation_elem = etree.fromstring(equation_xml.encode('utf-8'))
        run = paragraph.add_run()
        run._element.append(equation_elem)
    except Exception as e:
        print(f'Lỗi chèn công thức: {e}')
        run = paragraph.add_run('[CÔNG THỨC]')
        try:
            run.font.color.rgb = RGBColor(255, 0, 0)
        except:
            pass


def apply_formatting(run, formatting):
    """Áp dụng formatting cho run - ABSOLUTELY SAFE"""
    if not formatting:
        return

    try:
        run.bold = formatting.get('bold', False)
        run.italic = formatting.get('italic', False)
        run.underline = formatting.get('underline', False)

        font_name = formatting.get('font_name')
        if font_name:
            try:
                run.font.name = font_name
            except:
                pass

        font_size_pt = formatting.get('font_size')
        if font_size_pt:
            try:
                run.font.size = Pt(font_size_pt)
            except:
                pass

        # ULTIMATE FIX cho font_color
        font_color = formatting.get('font_color')
        if font_color:
            try:
                normalized = safe_normalize_rgb(font_color)
                if normalized:
                    run.font.color.rgb = RGBColor(normalized[0], normalized[1], normalized[2])
            except:
                pass

    except Exception as e:
        print(f'Lỗi apply formatting: {e}')


def write_elements_to_paragraph(p, elements, prefix=None):
    """Ghi các elements vào paragraph"""
    if prefix:
        run = p.add_run(prefix)
        run.bold = True

    text_buffer = ''
    first_real_text_element_index = -1
    last_text_run = None

    for i, el in enumerate(elements):
        if el.type == 'text':
            if first_real_text_element_index == -1 and el.data.strip():
                first_real_text_element_index = i

            if first_real_text_element_index == -1 or el.data == '\n':
                continue

            text_buffer += el.data

            regex_pattern = r'^\s*((Câu\s*\d+\.?)|(Bài\s*\d+\.?)|([A-D]\.\s*\*?)|([a-d]\))|Lời giải:)\s*'

            text_after_sub = re.sub(regex_pattern, '', text_buffer)

            if text_after_sub != text_buffer:
                for j in range(first_real_text_element_index, i + 1):
                    current_el = elements[j]
                    if current_el.type == 'text':
                        if j == i:
                            current_el.data = text_after_sub
                        else:
                            current_el.data = ''
            else:
                if text_buffer.strip():
                    break
        elif el.type in ['image', 'equation']:
            if first_real_text_element_index == -1:
                continue
            break

    for el in elements:
        if el.type == 'text':
            if el.data:
                text = el.data
                text = re.sub(r'\n\s*\n+', '\n', text)
                text = re.sub(r' +', ' ', text)

                if text.strip():
                    run = p.add_run(text)
                    apply_formatting(run, el.formatting)
                    last_text_run = run

        elif el.type == 'equation':
            if last_text_run is not None:
                p.add_run(' ')
            insert_equation_xml(p, el.data)
            p.add_run(' ')

        elif el.type == 'image':
            try:
                if last_text_run is not None:
                    p.add_run('\n')

                run = p.add_run()
                pic_width = Emu(el.width) if el.width else None
                pic_height = Emu(el.height) if el.height else None
                image_data = el.data

                if isinstance(image_data, str):
                    continue

                if not image_data:
                    continue

                run.add_picture(io.BytesIO(image_data), width=pic_width, height=pic_height)
                p.add_run('\n')

            except Exception as e:
                print(f'Không thể ghi hình ảnh: {e}')


def write_element_list_as_paragraphs(doc, element_list, prefix=None):
    """Ghi danh sách elements thành paragraph"""
    if not element_list:
        return None

    p = doc.add_paragraph()
    write_elements_to_paragraph(p, element_list, prefix=prefix)
    return p


def write_new_doc(filepath, parts, write_solution=False, exam_code=None):
    """Tạo file Word mới"""
    doc = docx.Document()

    if exam_code:
        p_code = doc.add_paragraph()
        run_code = p_code.add_run(f'Mã đề: {exam_code}')
        run_code.bold = True
        run_code.font.size = Pt(14)
        p_code.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()

    for part in parts:
        p_title = write_element_list_as_paragraphs(doc, part.title_elements)
        if p_title:
            p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p_title.runs:
                run.bold = True
                run.font.size = Pt(13)

        if part.intro_elements:
            write_element_list_as_paragraphs(doc, part.intro_elements)

        for idx, q in enumerate(part.questions):
            prefix = f"Bài {idx + 1}. " if part.type == 4 else f"Câu {idx + 1}. "

            write_element_list_as_paragraphs(doc, q.question_body, prefix=prefix)

            if part.type == 1:
                new_labels = ['A', 'B', 'C', 'D']
                # Ensure we don't exceed available labels
                num_options = min(len(q.options), len(new_labels))
                for i in range(num_options):
                    if i >= len(q.options):
                        break
                    opt = q.options[i]
                    write_element_list_as_paragraphs(doc, opt.content, prefix=f'{new_labels[i]}. ')

            elif part.type == 2:
                new_labels = ['a', 'b', 'c', 'd']
                # Ensure we don't exceed available labels
                num_options = min(len(q.options), len(new_labels))
                
                for i in range(num_options):
                    if i >= len(q.options):
                        break
                        
                    opt = q.options[i]
                    p_opt = doc.add_paragraph()
                    run_label = p_opt.add_run(f'{new_labels[i]}) ')
                    run_label.bold = True

                    for el in opt.content:
                        if el.type == 'text':
                            text = el.data
                            text = re.sub(r'[a-d]\)\s*', '', text)
                            text = text.replace('\n', ' ').strip()
                            text = re.sub(r'\s+', ' ', text)

                            if text:
                                run = p_opt.add_run(text)
                                apply_formatting(run, el.formatting)

                        elif el.type == 'equation':
                            insert_equation_xml(p_opt, el.data)

                        elif el.type == 'image':
                            try:
                                run = p_opt.add_run()
                                pic_width = Emu(el.width) if el.width else None
                                pic_height = Emu(el.height) if el.height else None
                                if el.data and not isinstance(el.data, str):
                                    run.add_picture(io.BytesIO(el.data), width=pic_width, height=pic_height)
                            except:
                                pass

            if write_solution:
                if q.solution_body:
                    write_element_list_as_paragraphs(doc, q.solution_body, prefix='Lời giải: ')

                if part.type != 4:
                    p_ans = doc.add_paragraph()
                    run_ans = p_ans.add_run(f'Đáp án đúng: {q.new_correct_answer}')
                    run_ans.bold = True

    doc.save(filepath)


def write_answer_key_doc(filepath, answer_key, exam_code=None):
    """Tạo file đáp án"""

    def set_cell_text(cell, text, bold=False, align='CENTER'):
        cell.text = str(text)
        p = cell.paragraphs[0]
        if align == 'CENTER':
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == 'LEFT':
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        if p.runs:
            p.runs[0].bold = bold
            p.runs[0].font.size = Pt(11)

    doc = docx.Document()

    if exam_code:
        p_code = doc.add_paragraph()
        run_code = p_code.add_run(f'Mã đề: {exam_code}')
        run_code.bold = True
        run_code.font.size = Pt(14)
        p_code.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()

    doc.add_heading('ĐÁP ÁN ĐỀ THI', level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    for part_title, keys in answer_key.items():
        part_title_lower = part_title.lower()

        if 'tự luận' in part_title_lower or 'phần 4' in part_title_lower:
            continue

        p_title = doc.add_paragraph()
        run_title = p_title.add_run(part_title)
        run_title.bold = True
        run_title.font.size = Pt(13)

        if not keys:
            doc.add_paragraph('\t(Không có câu hỏi)')
            doc.add_paragraph()
            continue

        parsed_keys = []
        for key in keys:
            match = re.match(r'^\s*([\d\w]+)\.\s*(.*)', key.strip())
            if match:
                parsed_keys.append((match.group(1), match.group(2)))

        if not parsed_keys:
            doc.add_paragraph('\t(Lỗi đọc đáp án)')
            doc.add_paragraph()
            continue

        try:
            if ';' in parsed_keys[0][1]:
                num_questions = len(parsed_keys)
                table = doc.add_table(rows=5, cols=num_questions + 1)
                table.style = 'Table Grid'

                set_cell_text(table.rows[0].cells[0], '', bold=True)

                for i, (num, ans) in enumerate(parsed_keys):
                    set_cell_text(table.rows[0].cells[i + 1], f'Câu {num}', bold=True)

                labels = ['a)', 'b)', 'c)', 'd)']
                for row_idx, label in enumerate(labels, start=1):
                    set_cell_text(table.rows[row_idx].cells[0], label, bold=True, align='LEFT')

                for col_idx, (num, ans) in enumerate(parsed_keys, start=1):
                    correct_answers = re.split(r'[;,\s]+', ans.strip())
                    correct_answers = [a.strip().lower() for a in correct_answers if a.strip()]

                    for row_idx, label in enumerate(['a', 'b', 'c', 'd'], start=1):
                        if label in correct_answers:
                            set_cell_text(table.rows[row_idx].cells[col_idx], 'Đ')
                        else:
                            set_cell_text(table.rows[row_idx].cells[col_idx], 'S')
            else:
                num_questions = len(parsed_keys)
                table = doc.add_table(rows=2, cols=num_questions + 1)
                table.style = 'Table Grid'

                set_cell_text(table.rows[0].cells[0], 'Câu', bold=True)
                set_cell_text(table.rows[1].cells[0], 'Đáp án', bold=True)

                for i, (num, ans) in enumerate(parsed_keys):
                    set_cell_text(table.rows[0].cells[i + 1], num, bold=True)
                    set_cell_text(table.rows[1].cells[i + 1], ans)

            doc.add_paragraph()
        except Exception as e:
            print(f'Lỗi vẽ bảng: {e}')
            for key in keys:
                doc.add_paragraph('\t' + key)

    doc.save(filepath)


def write_excel_summary(filepath, all_answer_data, parts_info):
    """Tạo file Excel tổng hợp đáp án"""
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print('Cần cài openpyxl: pip install openpyxl')
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Tổng hợp đáp án'

    max_questions = {1: 0, 2: 0, 3: 0, 4: 0}
    for data in all_answer_data:
        for part_key, answers in data['answer_key'].items():
            if 'phần 1' in part_key.lower() or 'trắc nghiệm (abcd)' in part_key.lower():
                max_questions[1] = max(max_questions[1], len(answers))
            elif 'phần 2' in part_key.lower() or 'đúng/sai' in part_key.lower():
                max_questions[2] = max(max_questions[2], len(answers))
            elif 'phần 3' in part_key.lower() or 'trả lời ngắn' in part_key.lower():
                max_questions[3] = max(max_questions[3], len(answers))

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    col = 1

    ws.cell(1, col, 'Mã đề')
    ws.cell(1, col).font = Font(bold=True, size=12, color='FFFFFF')
    ws.cell(1, col).fill = PatternFill(start_color='000000', end_color='000000', fill_type='solid')
    ws.cell(1, col).alignment = Alignment(horizontal='center', vertical='center')
    ws.cell(1, col).border = thin_border
    col += 1

    type_names = {
        1: 'TRẮC NGHIỆM (ABCD)',
        2: 'ĐÚNG/SAI (abcd)',
        3: 'TRẢ LỜI NGẮN'
    }

    colors = {
        1: '366092',
        2: '4472C4',
        3: '70AD47'
    }

    for q_type in [1, 2, 3]:
        if max_questions[q_type] > 0:
            start_col = col
            end_col = col + max_questions[q_type] - 1

            if start_col == end_col:
                ws.cell(1, start_col, type_names[q_type])
                ws.cell(1, start_col).font = Font(bold=True, size=11, color='FFFFFF')
                ws.cell(1, start_col).fill = PatternFill(start_color=colors[q_type], end_color=colors[q_type],
                                                         fill_type='solid')
                ws.cell(1, start_col).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(1, start_col).border = thin_border
            else:
                ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=end_col)
                ws.cell(1, start_col, type_names[q_type])
                ws.cell(1, start_col).font = Font(bold=True, size=11, color='FFFFFF')
                ws.cell(1, start_col).fill = PatternFill(start_color=colors[q_type], end_color=colors[q_type],
                                                         fill_type='solid')
                ws.cell(1, start_col).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(1, start_col).border = thin_border

            for i in range(max_questions[q_type]):
                ws.cell(2, col, f'Câu {i + 1}')
                ws.cell(2, col).font = Font(bold=True, size=10)
                ws.cell(2, col).fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
                ws.cell(2, col).alignment = Alignment(horizontal='center', vertical='center')
                ws.cell(2, col).border = thin_border
                col += 1

    row_idx = 3
    for data in all_answer_data:
        col = 1

        ws.cell(row_idx, col, data['exam_code'])
        ws.cell(row_idx, col).alignment = Alignment(horizontal='center', vertical='center')
        ws.cell(row_idx, col).font = Font(bold=True)
        ws.cell(row_idx, col).border = thin_border
        col += 1

        answers_by_type = {1: [], 2: [], 3: []}

        for part_key, answers in data['answer_key'].items():
            parsed_answers = []
            for ans_str in answers:
                match = re.match(r'^\s*(\d+)\.\s*(.*)', ans_str.strip())
                if match:
                    parsed_answers.append(match.group(2))

            if 'phần 1' in part_key.lower() or 'trắc nghiệm (abcd)' in part_key.lower():
                answers_by_type[1] = parsed_answers
            elif 'phần 2' in part_key.lower() or 'đúng/sai' in part_key.lower():
                answers_by_type[2] = parsed_answers
            elif 'phần 3' in part_key.lower() or 'trả lời ngắn' in part_key.lower():
                answers_by_type[3] = parsed_answers

        for q_type in [1, 2, 3]:
            for i in range(max_questions[q_type]):
                if i < len(answers_by_type[q_type]):
                    answer = answers_by_type[q_type][i]
                    ws.cell(row_idx, col, answer)
                    ws.cell(row_idx, col).alignment = Alignment(horizontal='center', vertical='center')
                    ws.cell(row_idx, col).border = thin_border
                else:
                    ws.cell(row_idx, col, '')
                    ws.cell(row_idx, col).border = thin_border
                col += 1

        row_idx += 1

    for col_num in range(1, col):
        column_letter = get_column_letter(col_num)
        max_length = 0

        for row_num in range(1, row_idx):
            try:
                cell_value = ws.cell(row_num, col_num).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            except:
                pass

        adjusted_width = min(max_length + 2, 15) if max_length > 0 else 10
        ws.column_dimensions[column_letter].width = adjusted_width

    ws.freeze_panes = 'B3'

    wb.save(filepath)


# ===================== MAIN APPLICATION =====================

class App(ctk.CTk):
    instance = None

    def __init__(self):
        super().__init__()
        App.instance = self

        DatabaseManager.init_db()

        self.title('Ngân Hàng Câu Hỏi & Tạo Đề Thi v3.5 ULTIMATE')
        self.geometry('650x780')
        ctk.set_appearance_mode('System')

        self.hybrid_scramble_parts_var = BooleanVar(value=False)
        self.hybrid_scramble_questions_var = BooleanVar(value=True)
        self.scramble_mcq_options_var = BooleanVar(value=True)
        self.scramble_tf_options_var = BooleanVar(value=True)
        self.hybrid_export_solution_file_var = BooleanVar(value=True)
        self.export_excel_var = BooleanVar(value=True)

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)

        self.label_title = ctk.CTkLabel(
            self,
            text='NGÂN HÀNG CÂU HỎI & TẠO ĐỀ THI',
            font=ctk.CTkFont(size=20, weight='bold')
        )
        self.label_title.grid(row=0, column=0, padx=20, pady=(20, 10))

        self.tab_view = ctk.CTkTabview(self, height=680)
        self.tab_view.grid(row=1, column=0, padx=20, pady=10, sticky='nsew')

        self.tab_create_hybrid = self.tab_view.add('Tạo đề từ ngân hàng')
        self.tab_manage_bank = self.tab_view.add('Quản lý ngân hàng')

        self.tab_view.set('Tạo đề từ ngân hàng')

        self.setup_create_hybrid_tab()
        self.setup_manage_bank_tab()

    def setup_manage_bank_tab(self):
        """Tab quản lý ngân hàng"""
        self.tab_manage_bank.grid_columnconfigure(0, weight=1)

        self.import_export_frame = ctk.CTkFrame(self.tab_manage_bank)
        self.import_export_frame.grid(row=0, column=0, padx=10, pady=10, sticky='ew')
        self.import_export_frame.grid_columnconfigure(0, weight=1)
        self.import_export_frame.grid_columnconfigure(1, weight=1)
        self.import_export_frame.grid_columnconfigure(2, weight=1)

        ctk.CTkLabel(
            self.import_export_frame,
            text='IMPORT/EXPORT NGÂN HÀNG',
            font=ctk.CTkFont(size=14, weight='bold')
        ).grid(row=0, column=0, columnspan=3, pady=(10, 5))

        self.btn_import_word = ctk.CTkButton(
            self.import_export_frame,
            text='📄 Import Word',
            command=self.start_import_thread,
            height=35
        )
        self.btn_import_word.grid(row=1, column=0, padx=5, pady=5, sticky='ew')

        self.btn_import_db = ctk.CTkButton(
            self.import_export_frame,
            text='📦 Import DB',
            command=self.import_database,
            height=35,
            fg_color='#2E7D32',
            hover_color='#1B5E20'
        )
        self.btn_import_db.grid(row=1, column=1, padx=5, pady=5, sticky='ew')

        self.btn_export_db = ctk.CTkButton(
            self.import_export_frame,
            text='💾 Export DB',
            command=self.export_database,
            height=35,
            fg_color='#1976D2',
            hover_color='#0D47A1'
        )
        self.btn_export_db.grid(row=1, column=2, padx=5, pady=5, sticky='ew')

        self.label_import_status = ctk.CTkLabel(
            self.import_export_frame,
            text='Sẵn sàng import/export',
            text_color='gray'
        )
        self.label_import_status.grid(row=2, column=0, columnspan=3, padx=10, pady=5)

        self.delete_frame = ctk.CTkFrame(self.tab_manage_bank, border_width=1, border_color='red')
        self.delete_frame.grid(row=1, column=0, padx=10, pady=10, sticky='ew')
        self.delete_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkLabel(
            self.delete_frame,
            text='QUẢN LÝ CHUYÊN ĐỀ',
            font=ctk.CTkFont(weight='bold', size=14),
            text_color='red'
        ).grid(row=0, column=0, pady=(10, 5))

        self.btn_clear_all = ctk.CTkButton(
            self.delete_frame,
            text='🗑️ XÓA TẤT CẢ',
            fg_color='red',
            hover_color='darkred',
            command=self.confirm_clear_all
        )
        self.btn_clear_all.grid(row=1, column=0, padx=10, pady=10, sticky='ew')

        self.topic_list_frame = ctk.CTkScrollableFrame(
            self.delete_frame,
            label_text='Danh sách chuyên đề',
            height=380
        )
        self.topic_list_frame.grid(row=2, column=0, padx=10, pady=10, sticky='ew')
        self.topic_list_frame.grid_columnconfigure(0, weight=1)

        self.display_delete_options()

    def setup_create_hybrid_tab(self):
        """Tab tạo đề từ ngân hàng"""
        self.tab_create_hybrid.grid_columnconfigure(0, weight=1)
        self.tab_create_hybrid.grid_rowconfigure(1, weight=1)

        self.refresh_frame = ctk.CTkFrame(self.tab_create_hybrid, fg_color='transparent')
        self.refresh_frame.grid(row=0, column=0, padx=10, pady=(5, 10), sticky='ew')

        self.btn_refresh_topics_hybrid = ctk.CTkButton(
            self.refresh_frame,
            text='🔄 LÀM MỚI DANH SÁCH',
            command=self.refresh_recipe_gui,
            height=35
        )
        self.btn_refresh_topics_hybrid.grid(row=0, column=0, padx=10, pady=5, sticky='ew')

        self.recipe_frame = ctk.CTkScrollableFrame(
            self.tab_create_hybrid,
            label_text='Chọn chuyên đề và số lượng câu hỏi',
            height=220
        )
        self.recipe_frame.grid(row=1, column=0, padx=10, pady=(0, 10), sticky='nsew')
        self.recipe_frame.grid_columnconfigure(1, weight=1)

        self.recipe_entries = {}
        self.refresh_recipe_gui()

        self.hybrid_options_frame = ctk.CTkFrame(self.tab_create_hybrid)
        self.hybrid_options_frame.grid(row=2, column=0, padx=10, pady=5, sticky='ew')
        self.hybrid_options_frame.grid_columnconfigure(0, weight=1)

        ctk.CTkCheckBox(
            self.hybrid_options_frame,
            text='Xáo trộn thứ tự các phần',
            variable=self.hybrid_scramble_parts_var
        ).grid(row=0, column=0, padx=10, pady=3, sticky='w')

        ctk.CTkCheckBox(
            self.hybrid_options_frame,
            text='Xáo trộn câu hỏi trong mỗi phần',
            variable=self.hybrid_scramble_questions_var
        ).grid(row=1, column=0, padx=10, pady=3, sticky='w')

        ctk.CTkCheckBox(
            self.hybrid_options_frame,
            text='Xáo trộn phương án trắc nghiệm',
            variable=self.scramble_mcq_options_var
        ).grid(row=2, column=0, padx=10, pady=3, sticky='w')

        ctk.CTkCheckBox(
            self.hybrid_options_frame,
            text='Xáo trộn phương án đúng/sai',
            variable=self.scramble_tf_options_var
        ).grid(row=3, column=0, padx=10, pady=3, sticky='w')

        ctk.CTkCheckBox(
            self.hybrid_options_frame,
            text='Xuất file lời giải riêng',
            variable=self.hybrid_export_solution_file_var
        ).grid(row=4, column=0, padx=10, pady=3, sticky='w')

        ctk.CTkCheckBox(
            self.hybrid_options_frame,
            text='📊 Xuất Excel tổng hợp đáp án',
            variable=self.export_excel_var,
            font=ctk.CTkFont(weight='bold')
        ).grid(row=5, column=0, padx=10, pady=3, sticky='w')

        self.hybrid_num_frame = ctk.CTkFrame(self.tab_create_hybrid)
        self.hybrid_num_frame.grid(row=3, column=0, padx=10, pady=5, sticky='ew')
        self.hybrid_num_frame.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(
            self.hybrid_num_frame,
            text='Số lượng đề:',
            font=ctk.CTkFont(size=13)
        ).grid(row=0, column=0, padx=10, pady=5, sticky='w')

        self.entry_num_copies_hybrid = ctk.CTkEntry(self.hybrid_num_frame, width=80)
        self.entry_num_copies_hybrid.insert(0, '4')
        self.entry_num_copies_hybrid.grid(row=0, column=1, padx=10, pady=5, sticky='w')

        ctk.CTkLabel(
            self.hybrid_num_frame,
            text='Mã đề (ngăn cách bởi ;):',
            font=ctk.CTkFont(size=13)
        ).grid(row=1, column=0, padx=10, pady=5, sticky='w')

        self.entry_custom_codes_hybrid = ctk.CTkEntry(
            self.hybrid_num_frame,
            placeholder_text='VD: 101; 102; 103'
        )
        self.entry_custom_codes_hybrid.grid(row=1, column=1, padx=10, pady=5, sticky='ew')

        self.hybrid_run_frame = ctk.CTkFrame(self.tab_create_hybrid, fg_color='transparent')
        self.hybrid_run_frame.grid(row=4, column=0, padx=10, pady=5, sticky='ew')

        self.btn_create_hybrid = ctk.CTkButton(
            self.hybrid_run_frame,
            text='🚀 BẮT ĐẦU TẠO ĐỀ',
            height=45,
            font=ctk.CTkFont(size=15, weight='bold'),
            fg_color='green',
            hover_color='darkgreen',
            command=self.start_hybrid_creation_thread
        )
        self.btn_create_hybrid.grid(row=0, column=0, padx=10, pady=10, sticky='ew')

        self.status_label_hybrid = ctk.CTkLabel(
            self.hybrid_run_frame,
            text='Sẵn sàng ...',
            text_color='blue'
        )
        self.status_label_hybrid.grid(row=1, column=0, padx=10, pady=5)

        self.progressbar_hybrid = ctk.CTkProgressBar(self.hybrid_run_frame, mode='indeterminate')

    def export_database(self):
        """Export database"""
        export_path = filedialog.asksaveasfilename(
            title='Lưu file Database',
            defaultextension='.db',
            filetypes=[('Database File', '*.db'), ('All Files', '*.*')]
        )

        if not export_path:
            return

        if DatabaseManager.export_database(export_path):
            self.update_import_status('✅ Export database thành công!', 'green')
            messagebox.showinfo('Thành công', f'Đã export database ra:\n{export_path}')
        else:
            self.update_import_status('❌ Lỗi export database', 'red')
            messagebox.showerror('Lỗi', 'Không thể export database')

    def import_database(self):
        """Import database"""
        import_path = filedialog.askopenfilename(
            title='Chọn file Database',
            filetypes=[('Database File', '*.db'), ('All Files', '*.*')]
        )

        if not import_path:
            return

        choice = messagebox.askyesnocancel(
            'Chọn phương thức Import',
            'YES = Thay thế toàn bộ database hiện tại\n'
            'NO = Gộp vào database hiện tại\n'
            'CANCEL = Hủy'
        )

        if choice is None:
            return

        self.update_import_status('Đang import...', 'orange')

        if choice:
            if DatabaseManager.import_database(import_path):
                self.update_import_status('✅ Thay thế database thành công!', 'green')
                messagebox.showinfo('Thành công', 'Đã thay thế database!')
                self.refresh_recipe_gui()
                self.display_delete_options()
            else:
                self.update_import_status('❌ Lỗi import database', 'red')
                messagebox.showerror('Lỗi', 'Không thể import database')
        else:
            if DatabaseManager.merge_database(import_path):
                self.update_import_status('✅ Gộp database thành công!', 'green')
                messagebox.showinfo('Thành công', 'Đã gộp database!')
                self.refresh_recipe_gui()
                self.display_delete_options()
            else:
                self.update_import_status('❌ Lỗi merge database', 'red')
                messagebox.showerror('Lỗi', 'Không thể merge database')

    def parse_custom_codes(self, codes_str):
        """Parse mã đề tùy chỉnh"""
        if not codes_str or not codes_str.strip():
            return None

        codes = [c.strip() for c in codes_str.split(';') if c.strip()]
        return codes if codes else None

    def update_import_status(self, text, color):
        """Cập nhật status import"""
        self.label_import_status.configure(text=text, text_color=color)

    def update_hybrid_status(self, text, color):
        """Cập nhật status tạo đề"""
        self.status_label_hybrid.configure(text=text, text_color=color)

    def start_import_thread(self):
        """Bắt đầu import file Word"""
        filepath = filedialog.askopenfilename(
            title='Chọn file Word để Import',
            filetypes=[('Word Documents', '*.docx')]
        )
        if not filepath:
            return

        topic_name = os.path.splitext(os.path.basename(filepath))[0]

        self.update_import_status('Đang phân tích file...', 'orange')

        threading.Thread(
            target=self.run_import_process,
            args=(filepath, topic_name),
            daemon=True
        ).start()

    def run_import_process(self, filepath, topic_name):
        """Xử lý import"""
        try:
            topic_id = DatabaseManager.get_or_create_topic(topic_name)
            parts = parse_word_doc(filepath)

            all_questions = []
            for part in parts:
                all_questions.extend(part.questions)

            if not all_questions:
                raise Exception('Không tìm thấy câu hỏi nào.')

            stats = {1: 0, 2: 0, 3: 0, 4: 0}
            for q in all_questions:
                q_blob = pickle.dumps(q)
                DatabaseManager.insert_question(topic_id, q.type, q_blob)
                stats[q.type] += 1

            type_names = {
                1: 'Trắc nghiệm',
                2: 'Đúng/Sai',
                3: 'Trả lời ngắn',
                4: 'Tự luận'
            }

            stats_text = ' | '.join([f'{type_names[t]}: {count}' for t, count in stats.items() if count > 0])

            self.update_import_status(
                f'✅ Import thành công {len(all_questions)} câu ({stats_text})',
                'green'
            )

            self.after(100, self.refresh_recipe_gui)
            self.after(100, self.display_delete_options)

        except Exception as e:
            tb_str = traceback.format_exc()
            error_msg = f'Lỗi import: {e}\n\n{tb_str}'
            self.update_import_status(f'❌ Lỗi: {e}', 'red')
            messagebox.showerror('Lỗi Import', error_msg)

    def confirm_clear_all(self):
        """Xác nhận xóa toàn bộ"""
        if messagebox.askyesno(
                'CẢNH BÁO',
                'Xóa TẤT CẢ câu hỏi và chuyên đề?\nHành động này không thể hoàn tác!'
        ):
            threading.Thread(target=self.run_clear_all_process, daemon=True).start()

    def run_clear_all_process(self):
        """Xóa toàn bộ database"""
        try:
            DatabaseManager.clear_all_questions()
            self.after(0, lambda: self.update_import_status('✅ Đã xóa toàn bộ', 'green'))
            self.after(100, self.refresh_recipe_gui)
            self.after(100, self.display_delete_options)
        except Exception as e:
            self.after(0, lambda: self.update_import_status(f'❌ Lỗi: {e}', 'red'))

    def confirm_delete_topic(self, topic_id, topic_name):
        """Xác nhận xóa chuyên đề"""
        if messagebox.askyesno(
                'Xác nhận',
                f'Xóa chuyên đề "{topic_name}" và TẤT CẢ câu hỏi?'
        ):
            threading.Thread(
                target=self.run_delete_topic_process,
                args=(topic_id, topic_name),
                daemon=True
            ).start()

    def run_delete_topic_process(self, topic_id, topic_name):
        """Xóa chuyên đề"""
        try:
            DatabaseManager.delete_topic_and_questions(topic_id)
            self.after(0, lambda: self.update_import_status(
                f'✅ Đã xóa "{topic_name}"',
                'green'
            ))
            self.after(100, self.refresh_recipe_gui)
            self.after(100, self.display_delete_options)
        except Exception as e:
            self.after(0, lambda: self.update_import_status(f'❌ Lỗi: {e}', 'red'))

    def confirm_edit_topic(self, topic_id, old_topic_name):
        """Sửa tên chuyên đề"""
        dialog = ctk.CTkInputDialog(
            text=f'Đổi tên chuyên đề "{old_topic_name}" thành:',
            title='Sửa tên chuyên đề'
        )
        new_topic_name = dialog.get_input()

        if not new_topic_name or not new_topic_name.strip() or new_topic_name == old_topic_name:
            return

        threading.Thread(
            target=self.run_edit_topic_process,
            args=(topic_id, old_topic_name, new_topic_name),
            daemon=True
        ).start()

    def run_edit_topic_process(self, topic_id, old_topic_name, new_topic_name):
        """Xử lý đổi tên"""
        try:
            DatabaseManager.rename_topic(topic_id, new_topic_name)
            self.after(0, lambda: self.update_import_status(
                f'✅ Đã đổi tên thành "{new_topic_name}"',
                'green'
            ))
            self.after(100, self.refresh_recipe_gui)
            self.after(100, self.display_delete_options)
        except Exception as e:
            self.after(0, lambda: self.update_import_status(f'❌ Lỗi: {e}', 'red'))

    def refresh_recipe_gui(self):
        """Làm mới danh sách chọn câu hỏi"""
        for widget in self.recipe_frame.winfo_children():
            widget.destroy()

        self.recipe_entries = {}
        all_topics = DatabaseManager.get_all_topics()
        sorted_topics = sorted(all_topics, key=lambda x: natural_sort_key(x[1]))

        if not sorted_topics:
            ctk.CTkLabel(
                self.recipe_frame,
                text='📭 Ngân hàng trống. Vui lòng import file ở tab "Quản lý ngân hàng".',
                font=ctk.CTkFont(size=12)
            ).grid(row=0, column=0, padx=10, pady=20)
            return

        row_idx = 0
        for topic_id, topic_name in sorted_topics:
            topic_label = ctk.CTkLabel(
                self.recipe_frame,
                text=f'📚 {topic_name}',
                font=ctk.CTkFont(weight='bold', size=13)
            )
            topic_label.grid(row=row_idx, column=0, columnspan=3, padx=10, pady=(10, 5), sticky='w')
            row_idx += 1

            for q_type in [1, 2, 3, 4]:
                count = DatabaseManager.count_questions_by_topic_and_type(topic_id, q_type)

                type_names = {
                    1: 'Trắc nghiệm (ABCD)',
                    2: 'Đúng/Sai (abcd)',
                    3: 'Trả lời ngắn',
                    4: 'Tự luận'
                }

                label = ctk.CTkLabel(
                    self.recipe_frame,
                    text=f'  {type_names[q_type]} (có {count}):',
                    anchor='w',
                    font=ctk.CTkFont(size=11)
                )
                label.grid(row=row_idx, column=0, padx=20, pady=2, sticky='w')

                entry = ctk.CTkEntry(self.recipe_frame, width=60)
                entry.grid(row=row_idx, column=1, padx=5, pady=2)

                self.recipe_entries[(topic_id, q_type)] = entry
                row_idx += 1

    def display_delete_options(self):
        """Hiển thị danh sách chuyên đề để xóa"""
        for widget in self.topic_list_frame.winfo_children():
            widget.destroy()

        self.topic_list_frame.grid_columnconfigure(0, weight=1)

        all_topics = DatabaseManager.get_all_topics()
        sorted_topics = sorted(all_topics, key=lambda x: natural_sort_key(x[1]))

        if not sorted_topics:
            ctk.CTkLabel(
                self.topic_list_frame,
                text='Ngân hàng trống'
            ).grid(row=0, column=0, columnspan=3, padx=10, pady=10)
            return

        for idx, (topic_id, topic_name) in enumerate(sorted_topics):
            total = sum(DatabaseManager.count_questions_by_topic_and_type(topic_id, t) for t in [1, 2, 3, 4])

            label = ctk.CTkLabel(
                self.topic_list_frame,
                text=f'{topic_name} ({total} câu)',
                anchor='w'
            )
            label.grid(row=idx, column=0, padx=10, pady=5, sticky='ew')

            btn_edit = ctk.CTkButton(
                self.topic_list_frame,
                text='✏️',
                width=40,
                command=lambda tid=topic_id, tname=topic_name: self.confirm_edit_topic(tid, tname)
            )
            btn_edit.grid(row=idx, column=1, padx=5, pady=5)

            btn_delete = ctk.CTkButton(
                self.topic_list_frame,
                text='🗑️',
                width=40,
                fg_color='red',
                hover_color='darkred',
                command=lambda tid=topic_id, tname=topic_name: self.confirm_delete_topic(tid, tname)
            )
            btn_delete.grid(row=idx, column=2, padx=5, pady=5)

    def start_hybrid_creation_thread(self):
        """Bắt đầu tạo đề"""
        recipe = {}

        try:
            for (topic_id, q_type), entry in self.recipe_entries.items():
                count_str = entry.get()
                if count_str:
                    count = int(count_str)
                    if count > 0:
                        recipe[(topic_id, q_type)] = count

            if not recipe:
                messagebox.showwarning('Thiếu thông tin', 'Chưa chọn số lượng câu hỏi.')
                return
        except ValueError:
            messagebox.showerror('Lỗi', 'Số lượng phải là số nguyên.')
            return

        custom_codes_str = self.entry_custom_codes_hybrid.get()
        custom_codes = self.parse_custom_codes(custom_codes_str)

        try:
            if custom_codes:
                num_copies = len(custom_codes)
            else:
                num_copies = int(self.entry_num_copies_hybrid.get())
                if num_copies <= 0:
                    raise ValueError()
        except ValueError:
            messagebox.showerror('Lỗi', 'Số lượng đề phải là số nguyên dương.')
            return

        self.btn_create_hybrid.configure(state='disabled')
        self.update_hybrid_status('Đang tạo đề...', 'orange')
        self.progressbar_hybrid.grid(row=2, column=0, padx=10, pady=10, sticky='ew')
        self.progressbar_hybrid.start()

        threading.Thread(
            target=self.run_hybrid_creation,
            args=(
                recipe,
                num_copies,
                self.hybrid_scramble_parts_var.get(),
                self.hybrid_scramble_questions_var.get(),
                self.scramble_mcq_options_var.get(),
                self.scramble_tf_options_var.get(),
                self.hybrid_export_solution_file_var.get(),
                self.export_excel_var.get(),
                custom_codes
            ),
            daemon=True
        ).start()

    def run_hybrid_creation(self, recipe, num_copies, scramble_parts_flag, scramble_questions_flag,
                            scramble_mcq_options_flag, scramble_tf_options_flag,
                            export_solution_file_flag, export_excel_flag, custom_codes=None):
        """Xử lý tạo đề"""
        try:
            base_dir = filedialog.askdirectory(title='Chọn nơi lưu đề thi')

            if not base_dir:
                self.update_hybrid_status('Đã hủy', 'gray')
                self.btn_create_hybrid.configure(state='normal')
                self.progressbar_hybrid.stop()
                self.progressbar_hybrid.grid_forget()
                return

            output_folder = os.path.join(base_dir, 'DE_THI_NGAN_HANG')

            try:
                if not os.path.exists(output_folder):
                    os.makedirs(output_folder)
            except PermissionError:
                desktop = os.path.join(os.path.expanduser('~'), 'Desktop')
                output_folder = os.path.join(desktop, 'DE_THI_NGAN_HANG')
                if not os.path.exists(output_folder):
                    os.makedirs(output_folder)

            all_answer_data = []
            parts_info = {}

            for i in range(1, num_copies + 1):
                self.update_hybrid_status(f'Đang tạo đề {i}/{num_copies}...', 'orange')

                if custom_codes:
                    exam_code = custom_codes[i - 1]
                else:
                    exam_code = f"{i:03d}"

                parts = []
                for (topic_id, q_type), count in recipe.items():
                    questions = DatabaseManager.get_random_questions(topic_id, q_type, count)

                    if not questions:
                        continue

                    deserialized = [pickle.loads(q_blob) for q_blob in questions]

                    topic_name = DatabaseManager.get_topic_name(topic_id)

                    type_names = {
                        1: 'Trắc nghiệm (ABCD)',
                        2: 'Đúng/Sai (abcd)',
                        3: 'Trả lời ngắn',
                        4: 'Tự luận'
                    }

                    part_title = f'Phần {q_type}: {type_names[q_type]} - {topic_name}'

                    part = Part(type=q_type, title_elements=[
                        ContentElement('text', part_title, {'bold': True, 'font_size': 13})
                    ])
                    part.questions = deserialized
                    parts.append(part)

                if not parts:
                    raise Exception('Không lấy được câu hỏi.')

                scrambled_parts, answer_key = scramble_data(
                    parts,
                    scramble_parts_flag,
                    scramble_questions_flag,
                    scramble_mcq_options_flag,
                    scramble_tf_options_flag
                )

                all_answer_data.append({
                    'exam_code': exam_code,
                    'answer_key': answer_key
                })

                output_doc_name = f'DE_{exam_code}.docx'
                output_doc_path = os.path.join(output_folder, output_doc_name)

                try:
                    write_new_doc(output_doc_path, scrambled_parts, write_solution=False, exam_code=exam_code)
                except PermissionError:
                    raise Exception(f'Không thể ghi file {output_doc_name}. Vui lòng đóng file nếu đang mở.')

                output_key_name = f'DAP_AN_{exam_code}.docx'
                output_key_path = os.path.join(output_folder, output_key_name)
                write_answer_key_doc(output_key_path, answer_key, exam_code=exam_code)

                if export_solution_file_flag:
                    output_solution_name = f'LOI_GIAI_{exam_code}.docx'
                    output_solution_path = os.path.join(output_folder, output_solution_name)
                    write_new_doc(output_solution_path, scrambled_parts, write_solution=True, exam_code=exam_code)

            if export_excel_flag and all_answer_data:
                self.update_hybrid_status('Đang tạo Excel tổng hợp...', 'orange')
                excel_path = os.path.join(output_folder, 'TONG_HOP_DAP_AN.xlsx')
                write_excel_summary(excel_path, all_answer_data, parts_info)

            self.update_hybrid_status(
                f'✅ Hoàn thành! Đã tạo {num_copies} đề',
                'green'
            )
            messagebox.showinfo('Hoàn thành', f'Đã tạo {num_copies} đề thi!\nLưu tại: {output_folder}')

            try:
                os.startfile(output_folder)
            except:
                pass

            self.btn_create_hybrid.configure(state='normal')
            self.progressbar_hybrid.stop()
            self.progressbar_hybrid.grid_forget()

        except Exception as e:
            tb_str = traceback.format_exc()
            error_msg = f'Lỗi: {e}\n\n{tb_str}'
            self.update_hybrid_status(f'❌ Lỗi: {e}', 'red')
            messagebox.showerror('Lỗi', error_msg)
            self.btn_create_hybrid.configure(state='normal')
            self.progressbar_hybrid.stop()
            self.progressbar_hybrid.grid_forget()


# ===================== MAIN =====================

def main():
    """Khởi chạy ứng dụng"""
    try:
        app = App()
        app.mainloop()
    except Exception as e:
        print(f'Lỗi khởi động: {e}')
        traceback.print_exc()


if __name__ == '__main__':
    try:
        import customtkinter
        import docx
        import lxml

        main()
    except ImportError:
        print('=' * 60)
        print('LỖI: THIẾU THƯ VIỆN')
        print('pip install customtkinter python-docx lxml openpyxl')
        print('=' * 60)
        input('Nhấn Enter để thoát...')
        exit()
